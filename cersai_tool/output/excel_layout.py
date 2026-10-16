"""The working and final Excel files of a batch.

  1 Working_<batch>.xlsx  made by step 1: every value the tool found.
  3 Final_<batch>.xlsx    made by step 2: the same, plus the values typed in
                          the required-details file. Step 3 reads this file,
                          including any correction made in it by hand, and
                          writes the .DAT file from it.

Sheet "CERSAI rows" is laid out like the demo file: the FH row, then for each
collateral its SI, IMM, BOR... and CHG rows, with a blank row between
collaterals. Values start in column B, one CERSAI field per column (sheet
"Columns" shows which field is in which column for each row type). Columns
AD-AF name the collateral, application and borrower of each row; they are not
part of the .DAT line. In the final file, column A joins the row with "|": a
preview of the .DAT line.
"""
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import settings
from ..model import BLANK, CONFLICT, INVALID, MISSING, NA, SRC_MANUAL, SRC_PDF
from ..rows import bor_row, chg_row, fh_row, imm_row, si_row
from ..rules.text import as_text

FIELDS_BY_TYPE = {"FH": fh_row.FIELDS, "SI": si_row.FIELDS, "IMM": imm_row.FIELDS, "BOR": bor_row.FIELDS,
                  "CHG": chg_row.FIELDS}
FIRST_VALUE_COLUMN = 2                     # column B
COLLATERAL_COLUMN, APPLICATION_COLUMN, BORROWER_COLUMN = 30, 31, 32     # AD, AE, AF
ROWS_SHEET = "CERSAI rows"

# five clearly different colours (black text stays readable on each)
FILLS = {
    "lms": "FF92D050",        # green: found in CKYC / COLLATERAL, fixed, or worked out
    "pdf": "FFB4A7FF",        # lavender: taken from the title report
    "manual": "FF5BC0EB",     # sky blue: typed in the required-details file
    "conflict": "FFFF9900",   # orange: LMS and title report disagree
    "problem": "FFFF3B3B",    # red: missing or invalid - must be fixed
}
LEGEND = [
    ("lms", "Found in CKYC / COLLATERAL, a fixed value, or worked out from them"),
    ("pdf", "Taken from the legal title report (the LMS did not have it)"),
    ("manual", "Typed in the required-details file"),
    ("conflict", "The LMS and the title report disagree: the value shown is the one used; check it"),
    ("problem", "Missing or invalid: must be filled before the .DAT file is made"),
]
ISSUE_HEADERS = ["Collateral ID", "Application number", "Row", "Borrower", "Field", "Level", "Problem", "Details"]
LEVEL_ORDER = {"hold": 0, "needs value": 1, "check": 2, "warning": 3, "info": 4}


def fill_key(cell):
    if cell.status in (MISSING, INVALID):
        return "problem"
    if cell.status == CONFLICT:
        return "conflict"
    if cell.status in (BLANK, NA):
        return None
    if cell.source == SRC_MANUAL:
        return "manual"
    if cell.source == SRC_PDF:
        return "pdf"
    return "lms"


def _fill(key):
    return PatternFill("solid", fgColor=FILLS[key])


def _comment_text(cell) -> str:
    lines = [cell.field]
    if cell.source:
        lines.append(f"From: {cell.source}")
    if cell.status not in ("ok",):
        lines.append(f"Status: {cell.status}")
    if cell.note:
        lines.append(cell.note)
    if cell.status == CONFLICT:
        lines.append(f"LMS: {cell.lms_value}")
        lines.append(f"Title report: {cell.pdf_value}")
    return "\n".join(lines)


def _borrower_name(borrower):
    return borrower["name"] if borrower else ""


def cell_problems(block) -> list:
    """Every cell of the block that needs a value or a check, as Issues-sheet rows."""
    out = []
    for kind, cells, borrower in block.rows:
        for cell in cells:
            if cell.status == MISSING:
                level, what = "needs value", "Missing"
            elif cell.status == INVALID:
                level, what = "needs value", "Invalid"
            elif cell.status == CONFLICT:
                level, what = "check", "LMS and title report differ"
            else:
                continue
            details = cell.note
            if cell.status == CONFLICT:
                details += f" | LMS: {cell.lms_value} | Title report: {cell.pdf_value} | used: {cell.value}"
            out.append([block.collateral_id, block.application_number, kind, _borrower_name(borrower),
                        cell.field, level, what, details])
    return out


def issue_rows(blocks, batch_issues) -> list:
    rows = []
    for issue in batch_issues:
        rows.append([issue.collateral_id, "", "", "", "", issue.level, issue.topic, issue.message])
    for block in blocks:
        block_rows = [[block.collateral_id, block.application_number, "", "", "", i.level, i.topic, i.message]
                      for i in block.issues()]
        block_rows += cell_problems(block)
        block_rows.sort(key=lambda r: LEVEL_ORDER.get(r[5], 9))
        rows.extend(block_rows)
    return rows


def write_batch_excel(path, fh_cells, blocks, batch_issues, *, final: bool, batch_name: str):
    book = Workbook()
    sheet = book.active
    sheet.title = ROWS_SHEET
    comments = settings.WRITE_FIELD_NAME_COMMENTS

    def put_row(r, kind, cells, collateral="", application="", borrower=""):
        for i, cell in enumerate(cells):
            target = sheet.cell(r, FIRST_VALUE_COLUMN + i)
            target.number_format = "@"
            target.value = cell.value if cell.value != "" else None
            key = fill_key(cell)
            if key:
                target.fill = _fill(key)
            if comments:
                note = Comment(_comment_text(cell), "CERSAI tool")
                note.width, note.height = 320, 140
                target.comment = note
        for column, value in ((COLLATERAL_COLUMN, collateral), (APPLICATION_COLUMN, application),
                              (BORROWER_COLUMN, borrower)):
            if value:
                sheet.cell(r, column, value).font = Font(color="FF808080")
        if final:
            parts = [f"{get_column_letter(FIRST_VALUE_COLUMN + i)}{r}" for i in range(len(cells))]
            sheet.cell(r, 1, "=" + '&"|"&'.join(parts)).font = Font(color="FF595959")

    r = 1
    put_row(r, "FH", fh_cells)
    r += 2
    for block in blocks:
        for kind, cells, borrower in block.rows:
            put_row(r, kind, cells, block.collateral_id, block.application_number,
                    borrower["neo_cif"] if borrower else "")
            r += 1
        r += 1

    sheet.column_dimensions["A"].width = 90 if final else 3
    sheet.column_dimensions["B"].width = 6
    for column in range(3, 29):
        sheet.column_dimensions[get_column_letter(column)].width = 15
    for column in (COLLATERAL_COLUMN, APPLICATION_COLUMN, BORROWER_COLUMN):
        sheet.column_dimensions[get_column_letter(column)].width = 15
    sheet.cell(1, COLLATERAL_COLUMN, "Collateral ID").font = Font(bold=True)
    sheet.cell(1, APPLICATION_COLUMN, "Application").font = Font(bold=True)
    sheet.cell(1, BORROWER_COLUMN, "Neo CIF ID").font = Font(bold=True)

    _write_columns_sheet(book.create_sheet("Columns"))
    _write_issues_sheet(book.create_sheet("Issues"), issue_rows(blocks, batch_issues))
    _write_read_me(book.create_sheet("Read me"), final, batch_name)
    book.save(path)


def _write_columns_sheet(sheet):
    sheet.cell(1, 1, "Which CERSAI field is in which column of sheet CERSAI rows").font = Font(bold=True)
    for r, (kind, fields) in enumerate(FIELDS_BY_TYPE.items(), start=3):
        sheet.cell(r, 1, kind).font = Font(bold=True)
        for i, name in enumerate(fields):
            sheet.cell(r, FIRST_VALUE_COLUMN + i, name)
    sheet.cell(2, 1, "Row")
    for i in range(27):
        sheet.cell(2, FIRST_VALUE_COLUMN + i, get_column_letter(FIRST_VALUE_COLUMN + i)).font = Font(bold=True)
    for column in range(1, 29):
        sheet.column_dimensions[get_column_letter(column)].width = 24 if column > 1 else 8


def _write_issues_sheet(sheet, rows):
    sheet.append(ISSUE_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFD9D9D9")
    for row in rows:
        sheet.append(row)
    for column, width in zip("ABCDEFGH", (15, 16, 6, 26, 26, 12, 30, 110)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = f"A1:H{len(rows) + 1}"


def _write_read_me(sheet, final, batch_name):
    lines = [f"CERSAI batch {batch_name} - {'final' if final else 'working'} file", ""]
    if final:
        lines += [
            "Made by step 2 from the working file and the required-details file.",
            "Step 3 turns this file into the .DAT file. You may correct values directly in sheet",
            "'CERSAI rows': step 3 uses the file exactly as it is when you run it. If you run step 2",
            "again later, this file is replaced (a copy of the edited file is kept as a backup).",
            "Do not insert or delete columns or rows inside a collateral's block.",
            "Column A previews each .DAT line.",
        ]
    else:
        lines += [
            "Made by step 1: every value the tool found in CKYC, COLLATERAL and the title reports.",
            "Nothing needs to be typed here: missing values go in the required-details file",
            "(2 Required details_<batch>.xlsx) in the same folder; step 2 then makes the final file.",
        ]
    lines += ["", "Every cell is text. Its comment names the CERSAI field, where the value came from, and",
              "any problem. Sheet 'Columns' lists the fields of each row type; sheet 'Issues' lists",
              "every problem found.", "", "Colours:"]
    for r, text in enumerate(lines, start=1):
        sheet.cell(r, 1, text)
    sheet.cell(1, 1).font = Font(bold=True, size=13)
    r = len(lines) + 1
    for key, text in LEGEND:
        sheet.cell(r, 1).fill = _fill(key)
        sheet.cell(r, 2, text)
        r += 1
    sheet.cell(r, 2, "No colour: optional value left empty, or a field that does not apply to this row")
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 100
    for row in sheet.iter_rows(min_row=1, max_row=len(lines)):
        row[0].alignment = Alignment(wrap_text=False)


def read_batch_excel(path) -> dict:
    """Reads sheet 'CERSAI rows' of a final Excel file.

    Returns {"fh": [values], "blocks": [{collateral_id, application_number,
    rows: [{"type", "values", "excel_row", "neo_cif"}]}], "problems": [text]}."""
    book = load_workbook(path, data_only=True)
    if ROWS_SHEET not in book.sheetnames:
        raise ValueError(f"{path.name} has no sheet '{ROWS_SHEET}'")
    sheet = book[ROWS_SHEET]
    fh, blocks, problems, current = None, [], [], None
    for number, values in enumerate(sheet.iter_rows(min_row=1, max_col=BORROWER_COLUMN, values_only=True),
                                    start=1):
        values = list(values) + [None] * (BORROWER_COLUMN - len(values))
        kind = as_text(values[FIRST_VALUE_COLUMN - 1]).upper()
        if not kind:
            continue
        if kind not in FIELDS_BY_TYPE:
            problems.append(f"row {number}: unknown row type '{kind}' in column B")
            continue
        count = len(FIELDS_BY_TYPE[kind])
        row_values = [as_text(v) for v in values[FIRST_VALUE_COLUMN - 1: FIRST_VALUE_COLUMN - 1 + count]]
        extra = [as_text(v) for v in values[FIRST_VALUE_COLUMN - 1 + count: COLLATERAL_COLUMN - 1] if as_text(v)]
        if extra:
            problems.append(f"row {number} ({kind}): values found after its last field: {extra}")
        collateral = as_text(values[COLLATERAL_COLUMN - 1])
        application = as_text(values[APPLICATION_COLUMN - 1])
        neo_cif = as_text(values[BORROWER_COLUMN - 1])
        if kind == "FH":
            fh = row_values
            continue
        if kind == "SI":
            current = {"collateral_id": collateral, "application_number": application, "rows": []}
            blocks.append(current)
        elif current is None:
            problems.append(f"row {number}: {kind} row before any SI row")
            continue
        current["rows"].append({"type": kind, "values": row_values, "excel_row": number, "neo_cif": neo_cif})
    book.close()
    return {"fh": fh, "blocks": blocks, "problems": problems}
