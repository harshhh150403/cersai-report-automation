"""Step 3: check the final Excel and write the .DAT file.

1. Picks the batch (the latest one with a final Excel, unless one is named).
2. Reads "3 Final_<batch>.xlsx" exactly as it is, hand corrections included.
3. Checks every SI (see output/checks.py). SIs with errors, third-party
   mortgages, held SIs and SIs already filed are left out and listed.
4. Writes the .DAT file(s) with the SIs that passed, and "4 Checks_<batch>.xlsx"
   listing what was included, what was left out and why, and any warnings.
5. Updates the Filed register: included SIs become Filed; the others are
   Pending with the reason, and come back in the next batch.
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .. import batch as batches
from ..output.checks import check_batch
from ..output.dat_writer import write_dat_files
from ..output.excel_layout import read_batch_excel
from ..sources.masters import Masters


def _sheet(book, title, headers, rows, widths):
    sheet = book.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFD9D9D9")
    for row in rows:
        sheet.append(row)
    for i, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + i)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def write_checks_report(path, folder, results, written, notes, moment):
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    included = [r for r in results if r.included]
    lines = [
        (f"CERSAI batch {folder.name} - checks before the .DAT file", True),
        (f"Checked on {moment.strftime('%Y-%m-%d %H:%M:%S')}", False),
        ("", False),
        (f"SIs in the final Excel: {len(results)}", False),
        (f"Included in the .DAT file: {len(included)}", False),
        (f"Left out: {len(results) - len(included)} (sheet 'Left out')", False),
        (f"Warnings: {sum(len(r.warnings) for r in results)} (sheet 'Warnings')", False),
        ("", False),
    ]
    lines += [(f".DAT file: {path_.name} ({len(ids)} SI(s))", False) for path_, ids in written] or \
        [("No .DAT file was written: no SI passed the checks.", False)]
    if notes:
        lines += [("", False), ("Notes", True)] + [(note, False) for note in notes]
    for r, (text, bold) in enumerate(lines, start=1):
        summary.cell(r, 1, text).font = Font(bold=bold, size=13 if r == 1 else 11)
    summary.column_dimensions["A"].width = 120

    file_of = {coll: path_.name for path_, ids in written for coll in ids}
    _sheet(book, "Included", ["Collateral ID", "Application number", "DAT file"],
           [[r.collateral_id, r.application_number, file_of.get(r.collateral_id, "")] for r in included],
           [16, 18, 50])
    left_rows = []
    for r in results:
        if r.included:
            continue
        if r.left_out_reason:
            left_rows.append([r.collateral_id, r.application_number, "", "", r.left_out_reason])
        for label, field, problem in r.errors:
            left_rows.append([r.collateral_id, r.application_number, label, field, problem])
    _sheet(book, "Left out", ["Collateral ID", "Application number", "Row", "Field", "Reason"], left_rows,
           [16, 18, 22, 30, 100])
    warn_rows = [[r.collateral_id, r.application_number, label, field, problem]
                 for r in results for label, field, problem in r.warnings]
    _sheet(book, "Warnings", ["Collateral ID", "Application number", "Row", "Field", "Warning"], warn_rows,
           [16, 18, 22, 30, 100])
    book.save(path)


def run(paths, batch_name=None, assume_yes=False, log=print, moment=None):
    moment = moment or datetime.now()
    folder = batches.pick_batch(paths.output_folder, batch_name, "final", assume_yes, log)
    files = batches.batch_files(folder)
    data = batches.load_data(folder)
    notes = []
    if files["required"].exists() and files["required"].stat().st_mtime > files["final"].stat().st_mtime + 1:
        notes.append("The required-details file was saved after the final Excel was made. If it has new "
                     "values, run step 2 again first; this run uses the final Excel as it is.")
    if data.get("final_hash") and batches.file_hash(files["final"]) != data["final_hash"]:
        notes.append("The final Excel was changed by hand after step 2: the changed values are used.")
    for note in notes:
        log(f"  note: {note}")

    log("Reading the CERSAI masters...")
    masters = Masters(paths.masters_file)
    log(f"Checking {files['final'].name}...")
    sheet = read_batch_excel(files["final"])
    notes += sheet["problems"]
    for problem in sheet["problems"]:
        log(f"  problem: {problem}")
    register = batches.load_register(paths.register_file)
    held = data.get("held", {})
    if isinstance(held, list):          # batches made before hold reasons were recorded
        held = {coll: "the title report may belong to another property (answer PDF_CHECK)" for coll in held}
    results = check_batch(sheet, masters, held, register, batches.FILED)
    for r in results:
        for label, field, text in data.get("notices", {}).get(r.collateral_id, []):
            r.warnings.append((label, field, f"{text} (see the Issues sheet of the final Excel)"))
    included = [r for r in results if r.included]

    written = write_dat_files(folder, [r.block for r in included], moment) if included else []
    checks_path = files["checks"]
    if checks_path.exists():
        batches.backup(checks_path)
    write_checks_report(checks_path, folder, results, written, notes, moment)

    today = moment.date().isoformat()
    for path, ids in written:
        for coll in ids:
            app = next(r.application_number for r in included if r.collateral_id == coll)
            register[coll] = {"Collateral ID": coll, "Application number": app, "Batch": folder.name,
                              "Status": batches.FILED, "DAT file": path.name, "Date": today, "Reason": ""}
    for r in results:
        if r.included or not r.collateral_id:
            continue
        entry = register.get(r.collateral_id)
        if entry and entry.get("Status") == batches.FILED:
            continue
        reason = r.left_out_reason or "; ".join(f"{label} {field}: {problem}".strip()
                                                for label, field, problem in r.errors[:3])
        register[r.collateral_id] = {"Collateral ID": r.collateral_id, "Application number": r.application_number,
                                     "Batch": folder.name, "Status": batches.PENDING, "DAT file": "",
                                     "Date": today, "Reason": reason}
    batches.save_register(paths.register_file, register)

    log(f"  included: {len(included)} of {len(results)} SI(s)")
    for path, ids in written:
        log(f"Wrote {path.name} ({len(ids)} SI(s))")
    if not written:
        log("No .DAT file written: no SI passed the checks.")
    left_out = [r for r in results if not r.included]
    for r in left_out[:15]:
        why = r.left_out_reason or f"{len(r.errors)} error(s), e.g. " + \
            " ".join(str(x) for x in r.errors[0] if x)
        log(f"  left out - {r.collateral_id}: {why}")
    if len(left_out) > 15:
        log(f"  ... and {len(left_out) - 15} more (sheet 'Left out' of {checks_path.name})")
    log(f"Wrote {checks_path.name} and updated {paths.register_file.name}")
    return written
