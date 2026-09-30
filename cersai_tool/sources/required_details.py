"""The required-details file: "2 Required details_<batch>.xlsx".

Step 1 writes it; a person fills it; step 2 reads it. It replaces the old
hand-made "Req details.xlsx".

Sheets
  Read me            how to fill the file.
  Borrowers          one row per borrower (application + Neo CIF ID) with a
                     value to fill. A value typed here is used on every SI of
                     that application. Gender and father's/mother's name are
                     never in the sources, so most individuals are listed.
  Collaterals        one row per collateral with a value to fill.
  Check and correct  what someone should look at: differences between the LMS
                     and the title report (both values side by side), title
                     reports that may belong to another property, owner names
                     to confirm, and empty rows to correct any other value.
  Lists              (hidden) the choices offered by the drop-downs.

Cells to fill are yellow and empty. Grey cells show the value the tool found:
leave them, or type over them to correct the value. Amber cells show a default
the tool used. Blue cells hold values typed for the same collateral or
borrower in an earlier batch. Typing BLANK leaves an optional value empty.

Reading: a value counts as typed when it differs from what the tool found on
its own, so grey cells left as they are change nothing. Values in "Check and
correct" are applied last, so they win over the other sheets.
"""
import re
from collections import OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..context import PDF_OK, WRONG_PDF
from ..model import (BLANK, BORROWER_KEY_OF_FIELD, CLEAR_WORD, CONFLICT, DEFAULT, INVALID, MISSING, NA,
                     NEEDS_INPUT, SRC_MANUAL, Overrides, is_clear)
from ..rows import bor_row, imm_row, si_row
from ..rules.owners import normalise_name
from ..rules.text import as_text

# (key, hint shown under the header, BOR fields the key fills)
BORROWER_INPUTS = [
    ("BORROWER_TYPE", "pick from the list", ["BORROWER_TYPE"]),
    ("PAN", "e.g. ABCDE1234F", ["INDIVIDUAL_PAN", "BORROWER_PAN"]),
    ("CKYC", "14 digits; needed when there is no PAN", ["INDIVIDUAL_CKYC_NUMBER", "BORROWER_CKYC_NUMBER"]),
    ("REG_NUMBER", "CIN (company) or LLPIN (LLP)", ["BORROWER_REGISTRATION_NUMBER"]),
    ("REG_DATE", "incorporation date, YYYY-MM-DD", ["BORROWER_REGISTRATION_DATE"]),
    ("GENDER", "M, F or T", ["GENDER"]),
    ("FATHER_MOTHER_NAME", "letters and spaces only", ["FATHER_MOTHER_NAME"]),
    ("DATE_OF_BIRTH", "YYYY-MM-DD", ["DATE_OF_BIRTH"]),
    ("ADDRESS", "without city, state and PIN", ["ADDRESS_LINE_1", "ADDRESS_LINE_2", "ADDRESS_LINE_3"]),
    ("PIN_CODE", "6 digits", ["PIN_CODE"]),
    ("CITY", "city name or code", ["CITY"]),
    ("DISTRICT", "district name or code", ["DISTRICT", "STATE"]),
]
COLLATERAL_INPUTS = [
    ("SI_TYPE_ID", "type of mortgage: pick from the list", ["SI_TYPE_ID"]),
    ("SI_TYPE_OTHERS", "only when SI type is 21", ["SI_TYPE_OTHERS"]),
    ("TYPE_OF_CHARGE", "e.g. FIRST CHARGE", ["TYPE_OF_CHARGE"]),
    ("ASSET_TYPE_ID", "pick from the list", ["ASSET_TYPE_ID"]),
    ("ASSET_TYPE_OTHERS", "only when asset type is 3", ["ASSET_TYPE_OTHERS"]),
    ("ASSET_SUB_TYPE_ID", "pick from the list", ["ASSET_SUB_TYPE_ID"]),
    ("ASSET_DESCRIPTION", "max 100 characters", ["ASSET_DESCRIPTION"]),
    ("SURVEY_NUMBER", "several: separate with ^", ["SURVEY_NUMBER"]),
    ("PLOT_NUMBER", "survey or plot number needed", ["PLOT_NUMBER"]),
    ("ASSET_AREA", "number, max 99999.99", ["ASSET_AREA"]),
    ("ASSET_AREA_UNIT", "pick from the list", ["ASSET_AREA_UNIT"]),
    ("HOUSE_NUMBER", "flat / unit / house no.", ["HOUSE_NUMBER"]),
    ("FLOOR_NUMBER", "-3 to 100; ground = 0", ["FLOOR_NUMBER"]),
    ("BUILDING_NAME", "wing / tower / building", ["BUILDING_NAME"]),
    ("PROJECT_NAME", "society / project (mandatory)", ["PROJECT_NAME"]),
    ("STREET_NAME", "road / street", ["STREET_NAME"]),
    ("LOCALITY", "area / village (mandatory)", ["LOCALITY"]),
    ("PIN_CODE", "6 digits", ["PIN_CODE"]),
    ("CITY", "city name or code", ["CITY"]),
    ("DISTRICT", "district name or code", ["DISTRICT", "STATE"]),
    ("OWNERS", "names separated by ;", []),
]
BORROWER_KEYS = [k for k, _, _ in BORROWER_INPUTS]
COLLATERAL_KEYS = [k for k, _, _ in COLLATERAL_INPUTS]
# fields whose need for a value is shown in another column of the same sheet
BORROWER_COVERED = {"ADDRESS_LINE_1": "ADDRESS", "ADDRESS_LINE_2": "ADDRESS", "ADDRESS_LINE_3": "ADDRESS",
                    "STATE": "DISTRICT"}
COLLATERAL_COVERED = {"STATE": "DISTRICT", "NUMBER_OF_ASSET_OWNERS": "OWNERS"}
CODE_KEYS = {"BORROWER_TYPE", "SI_TYPE_ID", "ASSET_TYPE_ID", "ASSET_SUB_TYPE_ID", "ASSET_AREA_UNIT", "CITY",
             "DISTRICT"}
SPECIAL_FIELDS = {"OWNERS", "PDF_CHECK", "OWNER_MATCH"}

BORROWER_INFO = ["Application number", "Neo CIF ID", "Borrower name", "Collateral IDs"]
COLLATERAL_INFO = ["Collateral ID", "Application number", "Property address (COLLATERAL)", "Title report"]
CHECK_HEADERS = ["Collateral ID", "Application number", "Borrower Neo CIF ID", "Borrower name", "Row", "Field",
                 "What to check", "LMS value", "Title report value", "Value the tool will use", "Your value"]
EMPTY_CHECK_ROWS = 30

FILL_NEEDED = "FFFFFF00"     # yellow
FILL_DEFAULT = "FFFFC000"    # amber
FILL_FOUND = "FFE7E6E6"      # grey
FILL_CONFLICT = "FFF4B084"   # orange
FILL_TYPED = "FF9BC2E6"      # blue
FILL_NA = "FFA6A6A6"         # dark grey
FILL_HEADER = "FFD9D9D9"


def _fill(colour):
    return PatternFill("solid", fgColor=colour)


# --------------------------------------------------------------------------
# Showing a value: code fields are shown as "code - name"
# --------------------------------------------------------------------------
def label(masters, key, code) -> str:
    if not code:
        return ""
    names = {
        "BORROWER_TYPE": lambda c: masters.borrower_types.get(c, ""),
        "SI_TYPE_ID": lambda c: masters.si_types.get(c, ("", ""))[0],
        "ASSET_TYPE_ID": lambda c: masters.asset_types.get(c, ("", ""))[1],
        "ASSET_SUB_TYPE_ID": lambda c: masters.asset_sub_types.get(c, ("", ""))[1],
        "ASSET_AREA_UNIT": lambda c: masters.area_units.get(c, ""),
        "CITY": lambda c: masters.cities.get(c, ("", ""))[0],
        "DISTRICT": lambda c: masters.districts.get(c, ("", ""))[0],
    }
    name = names[key](code) if key in names else ""
    return f"{code} - {name}" if name else code


def normalise(key, value) -> str:
    """The value as the tool reads it: a code field keeps only the code of
    "code - name"; owner lists get uniform separators."""
    text = as_text(value)
    if key in CODE_KEYS:
        m = re.match(r"^\s*([A-Za-z0-9]+)\s+-\s", text)
        if m:
            text = m.group(1)
        return text.upper() if key == "BORROWER_TYPE" else text
    if key == "OWNERS":
        return "; ".join(part.strip() for part in text.split(";") if part.strip())
    return text


def _pick(cells, fields):
    """The cell that decides a key: the first of its fields that applies."""
    for name in fields:
        cell = cells.get(name)
        if cell is not None and cell.status != NA:
            return cell
    return cells.get(fields[0])


def shown_value(masters, key, fields, cells, owners=None):
    """(text, cell) the file shows for one key of a borrower or collateral."""
    if key == "OWNERS":
        return "; ".join(owners.owners) if owners else "", cells.get("NUMBER_OF_ASSET_OWNERS")
    cell = _pick(cells, fields)
    if cell is None:
        return "", None
    if key == "ADDRESS":
        lines = [cells[f].value for f in fields if f in cells and cells[f].value]
        return " ".join(lines), cell
    if key == "DISTRICT" and cells.get("STATE") is not None and cells["STATE"].status in (MISSING, INVALID) \
            and cell.status not in (MISSING, INVALID):
        return label(masters, key, cell.value), cells["STATE"]
    return (label(masters, key, cell.value) if key in CODE_KEYS else cell.value), cell


def _style(cell, typed_before) -> str:
    if cell is not None and cell.status in (MISSING, INVALID):
        return FILL_NEEDED
    if typed_before:
        return FILL_TYPED
    if cell is None:
        return FILL_FOUND
    if cell.status == DEFAULT:
        return FILL_DEFAULT
    if cell.status == CONFLICT:
        return FILL_CONFLICT
    if cell.status == NA:
        return FILL_NA
    if cell.status == BLANK:
        return None
    if cell.source == SRC_MANUAL:
        return FILL_TYPED
    return FILL_FOUND


def _note(cell) -> str:
    if cell is None:
        return ""
    status = {MISSING: "Needed", INVALID: "Invalid", DEFAULT: "Default used - confirm or type over it",
              CONFLICT: "LMS and title report differ - see sheet Check and correct"}.get(cell.status, "")
    lines = [line for line in (status, cell.note, f"From: {cell.source}" if cell.source else "") if line]
    if cell.status == CONFLICT:
        lines += [f"LMS: {cell.lms_value}", f"Title report: {cell.pdf_value}"]
    return "\n".join(lines)


# --------------------------------------------------------------------------
# Collecting what goes in the file
# --------------------------------------------------------------------------
def collect_borrowers(blocks) -> OrderedDict:
    """(application, Neo CIF) -> {borrower, cells, collaterals}. A borrower's
    values are the same on every SI of the application, so the first is used."""
    out = OrderedDict()
    for block in blocks:
        for kind, cells, borrower in block.rows:
            if kind != "BOR":
                continue
            key = (block.application_number, borrower["neo_cif"])
            entry = out.setdefault(key, {"borrower": borrower, "cells": {c.field: c for c in cells},
                                         "collaterals": [], "block": block})
            entry["collaterals"].append(block.collateral_id)
    return out


def collateral_cells(block) -> dict:
    cells = {}
    for kind, row_cells, _ in block.rows:
        if kind in ("SI", "IMM"):
            for cell in row_cells:
                cells.setdefault(cell.field, cell)
    return cells


def borrower_displays(masters, entry) -> dict:
    return {key: shown_value(masters, key, fields, entry["cells"])[0] for key, _, fields in BORROWER_INPUTS}


def collateral_displays(masters, block) -> dict:
    cells = collateral_cells(block)
    return {key: shown_value(masters, key, fields, cells, block.context.owners)[0]
            for key, _, fields in COLLATERAL_INPUTS}


# --------------------------------------------------------------------------
# Writing
# --------------------------------------------------------------------------
def write_required_details(path, blocks, carried: Overrides, masters, batch_name: str):
    book = Workbook()
    read_me = book.active
    read_me.title = "Read me"
    lists = _write_lists(book, masters, blocks)
    counts = {}
    counts["borrowers"] = _write_borrowers(book.create_sheet("Borrowers"), blocks, carried, masters, lists)
    counts["collaterals"] = _write_collaterals(book.create_sheet("Collaterals"), blocks, carried, masters, lists)
    counts["checks"] = _write_checks(book.create_sheet("Check and correct"), blocks, carried, masters, lists)
    _write_read_me(read_me, batch_name, counts)
    book.move_sheet("Lists", offset=len(book.sheetnames) - 1 - book.sheetnames.index("Lists"))
    book.save(path)
    return counts


def _write_lists(book, masters, blocks) -> dict:
    sheet = book.create_sheet("Lists")
    sheet.sheet_state = "hidden"
    columns = OrderedDict()
    columns["BORROWER_TYPE"] = [label(masters, "BORROWER_TYPE", c) for c in masters.borrower_types]
    columns["SI_TYPE_ID"] = [label(masters, "SI_TYPE_ID", c) for c in sorted(masters.immovable_si_types, key=int)]
    columns["ASSET_TYPE_ID"] = [label(masters, "ASSET_TYPE_ID", c) for c in ("1", "2", "3")]
    columns["ASSET_SUB_TYPE_ID"] = [f"{code} - {name} ({masters.asset_types.get(kind, ('', ''))[1]})"
                                    for code, (kind, name) in sorted(masters.asset_sub_types.items(),
                                                                     key=lambda x: int(x[0]))
                                    if kind in ("1", "2", "3")]
    columns["ASSET_AREA_UNIT"] = [label(masters, "ASSET_AREA_UNIT", c) for c in sorted(masters.area_units, key=int)]
    columns["GENDER"] = ["M", "F", "T"]
    columns["PDF_CHECK"] = [PDF_OK, WRONG_PDF]
    columns["YES_NO"] = ["Y", "N"]
    columns["ROW"] = ["SI", "IMM", "BOR"]
    fields = []
    for name in si_row.FIELDS + imm_row.FIELDS + bor_row.FIELDS + ["OWNERS"]:
        if name not in fields and name not in ("ROW_TYPE", "SERIAL_NUMBER"):
            fields.append(name)
    columns["FIELD"] = fields
    columns["COLLATERAL"] = [b.collateral_id for b in blocks]
    ranges = {}
    for i, (name, values) in enumerate(columns.items(), start=1):
        letter = get_column_letter(i)
        sheet.cell(1, i, name).font = Font(bold=True)
        for r, value in enumerate(values, start=2):
            sheet.cell(r, i, value)
        ranges[name] = f"Lists!${letter}$2:${letter}${max(2, len(values) + 1)}"
    return ranges


def _dropdown(sheet, formula, cells: str):
    rule = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
    sheet.add_data_validation(rule)
    for part in cells.split():
        rule.add(part)


def _header(sheet, info, inputs, widths):
    for i, name in enumerate(info + [k for k, _, _ in inputs], start=1):
        cell = sheet.cell(1, i, name)
        cell.font = Font(bold=True)
        cell.fill = _fill(FILL_HEADER)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, (_, hint, _) in enumerate(inputs, start=len(info) + 1):
        cell = sheet.cell(2, i, hint)
        cell.font = Font(italic=True, color="FF595959", size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.row_dimensions[2].height = 30


def _put_input(sheet, r, column, text, cell, typed_before):
    target = sheet.cell(r, column, text if text != "" else None)
    target.number_format = "@"
    colour = _style(cell, typed_before)
    if colour:
        target.fill = _fill(colour)
    note = _note(cell)
    if typed_before:
        note = "Typed in an earlier batch." + (f"\n{note}" if cell is not None and cell.status in (MISSING, INVALID)
                                               else "")
    if note:
        comment = Comment(note, "CERSAI tool")
        comment.width, comment.height = 300, 120
        target.comment = comment


def _carried_borrower(carried, app, cif) -> dict:
    return {key: value for (a, c, key), value in carried.borrower.items() if a == app and c == cif}


def _carried_collateral(carried, coll) -> dict:
    return {key: value for (c, key), value in carried.collateral.items() if c == coll}


def _write_borrowers(sheet, blocks, carried, masters, lists) -> int:
    widths = [15, 14, 28, 16] + [20, 13, 16, 22, 14, 10, 24, 13, 50, 10, 24, 24]
    _header(sheet, BORROWER_INFO, BORROWER_INPUTS, widths)
    r = 3
    for (app, cif), entry in collect_borrowers(blocks).items():
        cells = entry["cells"]
        typed = _carried_borrower(carried, app, cif)
        needs = any(c.status in NEEDS_INPUT for f, c in cells.items() if f != "IS_ASSET_OWNER")
        if not needs and not any(k in BORROWER_KEYS for k in typed):
            continue
        for i, value in enumerate([app, cif, entry["borrower"]["name"], ", ".join(entry["collaterals"])], 1):
            sheet.cell(r, i, value)
        for i, (key, _, fields) in enumerate(BORROWER_INPUTS, start=len(BORROWER_INFO) + 1):
            text, cell = shown_value(masters, key, fields, cells)
            typed_before = key in typed
            if typed_before:
                text = CLEAR_WORD if is_clear(typed[key]) else (
                    label(masters, key, normalise(key, typed[key])) if key in CODE_KEYS else as_text(typed[key]))
            _put_input(sheet, r, i, text, cell, typed_before)
        r += 1
    last = max(r, 200)
    first = len(BORROWER_INFO) + 1
    _dropdown(sheet, lists["BORROWER_TYPE"], f"{get_column_letter(first)}3:{get_column_letter(first)}{last}")
    gender = first + BORROWER_KEYS.index("GENDER")
    _dropdown(sheet, lists["GENDER"], f"{get_column_letter(gender)}3:{get_column_letter(gender)}{last}")
    sheet.freeze_panes = "D3"
    return r - 3


def _write_collaterals(sheet, blocks, carried, masters, lists) -> int:
    widths = [15, 15, 60, 30] + [24, 18, 16, 20, 18, 24, 40, 18, 14, 11, 16, 14, 11, 24, 30, 24, 24, 10, 24,
                                 24, 40]
    _header(sheet, COLLATERAL_INFO, COLLATERAL_INPUTS, widths)
    r = 3
    for block in blocks:
        cells = collateral_cells(block)
        typed = _carried_collateral(carried, block.collateral_id)
        owners = block.context.owners
        needs = any(c.status in NEEDS_INPUT for c in cells.values()) or owners.status == "missing"
        if not needs and not any(k in COLLATERAL_KEYS for k in typed):
            continue
        report = block.context.raw_report
        report_text = "none found" if not report else (
            report["file"] if report.get("status") == "ok" else f"{report['file']} (no readable text)")
        info = [block.collateral_id, block.application_number, block.context.bundle["property_address"], report_text]
        for i, value in enumerate(info, start=1):
            sheet.cell(r, i, value)
        for i, (key, _, fields) in enumerate(COLLATERAL_INPUTS, start=len(COLLATERAL_INFO) + 1):
            text, cell = shown_value(masters, key, fields, cells, owners)
            typed_before = key in typed
            if typed_before:
                text = CLEAR_WORD if is_clear(typed[key]) else (
                    label(masters, key, normalise(key, typed[key])) if key in CODE_KEYS else as_text(typed[key]))
            if key == "OWNERS" and owners.status == "missing" and not typed_before:
                target = sheet.cell(r, i)
                target.number_format = "@"
                target.fill = _fill(FILL_NEEDED)
                target.comment = Comment("Needed: no owners in COLLATERAL or the title report. "
                                         "Type the owners' names separated by ;", "CERSAI tool")
                continue
            _put_input(sheet, r, i, text, cell, typed_before)
        r += 1
    last = max(r, 200)
    first = len(COLLATERAL_INFO) + 1
    for key in ("SI_TYPE_ID", "ASSET_TYPE_ID", "ASSET_SUB_TYPE_ID", "ASSET_AREA_UNIT"):
        column = get_column_letter(first + COLLATERAL_KEYS.index(key))
        _dropdown(sheet, lists[key], f"{column}3:{column}{last}")
    sheet.freeze_panes = "B3"
    return r - 3


def owner_match_groups(blocks) -> OrderedDict:
    """Owner names that only probably match a borrower, grouped per
    application, borrower and owner name: the same owner on several
    collaterals of one application is asked about once.
    (application, Neo CIF, owner) -> [Collateral IDs]"""
    groups = OrderedDict()
    for block in blocks:
        for link in block.context.owner_links:
            if link.probable:
                key = (block.application_number, link.neo_cif, normalise_name(link.owner))
                groups.setdefault(key, [])
                if block.collateral_id not in groups[key]:
                    groups[key].append(block.collateral_id)
    return groups


def _check_rows(blocks, carried, masters) -> list:
    """Rows of the Check and correct sheet: [values..., your value, typed before, choices]."""
    rows = []
    borrowers_done = set()
    match_groups = owner_match_groups(blocks)
    for block in blocks:
        c = block.context
        coll, app = block.collateral_id, block.application_number
        typed = _carried_collateral(carried, coll)
        placed = set()

        def add(row_type, field, what, lms="", pdf="", used="", neo_cif="", name="", your="", choices=None):
            rows.append({"values": [coll, app, neo_cif, name, row_type, field, what, lms, pdf, used],
                         "your": your, "typed_before": bool(your), "choices": choices})

        # a title report that may belong to another property
        answer = typed.get("PDF_CHECK", "")
        if c.right_pdf["held"] or answer:
            report = c.candidate_report or {}
            failed = "; ".join(f"{n}: {d}" for n, passed, d in c.right_pdf["checks"] if not passed)
            sibling = c.right_pdf.get("sibling")
            what = (f"seems to be the report of {sibling}, another collateral of this application ({failed}): "
                    "if the two files were swapped, rename them and run step 1 again. Otherwise type"
                    if sibling else f"may belong to another property ({failed}). Type")
            add("SI", "PDF_CHECK",
                f"The title report {report.get('file', '')} {what} {PDF_OK} to use it, or {WRONG_PDF} to "
                "ignore it. Until then this SI is held.",
                f"borrowers: {', '.join(b['name'] for b in c.borrowers)}; owners: "
                f"{'; '.join(c.bundle['lms_owners'])}; address: {c.bundle['property_address']}",
                f"applicant: {report.get('loan_applicant', '')}; owners: "
                f"{'; '.join(report.get('owners_item4') or report.get('owners_header') or [])}; "
                f"unit: {', '.join(report.get('details', {}).get('unit_numbers') or [])}",
                "held" if c.right_pdf["held"] else "", your=answer, choices="PDF_CHECK")
            placed.add("PDF_CHECK")

        # LMS and title report disagree
        for kind, cells, _ in block.rows:
            if kind not in ("SI", "IMM"):
                continue
            for cell in cells:
                if cell.status == CONFLICT and cell.field != "NUMBER_OF_ASSET_OWNERS":
                    add(kind, cell.field, cell.note, cell.lms_value, cell.pdf_value, cell.value)

        # owners
        owners = c.owners
        messages = []
        if owners.status == "conflict":
            messages.append("The LMS and the title report name different owners; the title report's owners are "
                            "used.")
        elif owners.status == "held":
            messages.append("The LMS and the title report name different owners, and the title report may be "
                            "for another property (see PDF_CHECK); the LMS owners are used for now.")
        elif owners.status == "not_verified":
            messages.append("The owners could not be checked against a title report (none usable).")
        elif owners.status == "from_title_report":
            messages.append("COLLATERAL has no owners; the title report's owners are used.")
        outsiders = c.non_borrower_owners()
        if outsiders:
            messages.append(f"Owner(s) {', '.join(outsiders)} are not borrowers: third-party mortgage. This SI "
                            "will be left out of the .DAT file (asset-owner rows are not supported). If they are "
                            "borrowers, type the owners using the borrowers' names.")
        if messages:
            add("SI", "OWNERS", " ".join(messages) + " To correct, type the owners separated by ;",
                "; ".join(owners.lms_owners), "; ".join(owners.pdf_owners), "; ".join(owners.owners))

        # owner names that match a borrower only probably
        answers = {cif: value for (cl, cif, key), value in carried.per_owner.items()
                   if cl == coll and key == "OWNER_MATCH"}
        by_cif = {b["neo_cif"]: b for b in c.borrowers}
        for link in c.owner_links:
            if link.probable:
                group = match_groups[(app, link.neo_cif, normalise_name(link.owner))]
                if group[0] != coll:
                    continue          # asked on the application's first collateral with this owner
                also = (f" The answer is used for every collateral of this application with this owner: "
                        f"{', '.join(group)}." if len(group) > 1 else "")
                add("BOR", "OWNER_MATCH",
                    f"Is owner '{link.owner}' the same person as borrower '{link.borrower_name}'? Type Y or N. "
                    f"Until it is answered the SI is held (the values shown assume Y).{also}",
                    "; ".join(owners.lms_owners), "; ".join(owners.pdf_owners), f"Y (name score {link.score})",
                    link.neo_cif, link.borrower_name, choices="YES_NO")
        for cif, value in answers.items():
            add("BOR", "OWNER_MATCH", "Answer typed in an earlier batch: is this owner the same person as the "
                "borrower?", "; ".join(owners.lms_owners), "; ".join(owners.pdf_owners), "",
                cif, by_cif.get(cif, {}).get("name", ""), your=value, choices="YES_NO")
        for (cl, cif, key), value in carried.per_owner.items():
            if cl == coll and key != "OWNER_MATCH":
                add("BOR", key, "Typed in an earlier batch", neo_cif=cif, name=by_cif.get(cif, {}).get("name", ""),
                    your=value)

        # values needed that have no column in the Collaterals sheet
        for kind, cells, borrower in block.rows:
            if kind in ("SI", "IMM"):
                for cell in cells:
                    if cell.status in NEEDS_INPUT and cell.field not in COLLATERAL_KEYS \
                            and cell.field not in COLLATERAL_COVERED:
                        add(kind, cell.field, f"{cell.status}: {cell.note}", used=cell.value)
            elif kind == "BOR":
                key = (app, borrower["neo_cif"])
                if key in borrowers_done:
                    continue
                for cell in cells:
                    short = BORROWER_KEY_OF_FIELD.get(cell.field, cell.field)
                    if cell.status in NEEDS_INPUT and short not in BORROWER_KEYS \
                            and cell.field not in BORROWER_COVERED:
                        add("BOR", cell.field, f"{cell.status}: {cell.note}", used=cell.value,
                            neo_cif=borrower["neo_cif"], name=borrower["name"])

        # values typed in an earlier batch that have no column elsewhere
        for key, value in typed.items():
            if key not in COLLATERAL_KEYS and key not in placed:
                row_type = "SI" if key in si_row.FIELDS else "IMM"
                add(row_type, key, "Typed in an earlier batch", your=value)
        for borrower in c.borrowers:
            key = (app, borrower["neo_cif"])
            if key in borrowers_done:
                continue
            borrowers_done.add(key)
            for short, value in _carried_borrower(carried, app, borrower["neo_cif"]).items():
                if short not in BORROWER_KEYS:
                    add("BOR", short, "Typed in an earlier batch", neo_cif=borrower["neo_cif"],
                        name=borrower["name"], your=value)
    return rows


def _write_checks(sheet, blocks, carried, masters, lists) -> int:
    sheet.append(CHECK_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _fill(FILL_HEADER)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    rows = _check_rows(blocks, carried, masters)
    choice_cells = {}
    for r, row in enumerate(rows, start=2):
        for i, value in enumerate(row["values"], start=1):
            cell = sheet.cell(r, i, value if value != "" else None)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        your = sheet.cell(r, len(CHECK_HEADERS), row["your"] or None)
        your.number_format = "@"
        your.fill = _fill(FILL_TYPED if row["typed_before"] else "FFFFF2CC")
        if row["choices"]:
            choice_cells.setdefault(row["choices"], []).append(your.coordinate)
    first_empty = len(rows) + 2
    for r in range(first_empty, first_empty + EMPTY_CHECK_ROWS):
        sheet.cell(r, len(CHECK_HEADERS)).number_format = "@"
        sheet.cell(r, len(CHECK_HEADERS)).fill = _fill("FFFFF2CC")
    last = first_empty + EMPTY_CHECK_ROWS - 1
    for choices, coordinates in choice_cells.items():
        _dropdown(sheet, lists[choices], " ".join(coordinates))
    _dropdown(sheet, lists["COLLATERAL"], f"A{first_empty}:A{last}")
    _dropdown(sheet, lists["ROW"], f"E{first_empty}:E{last}")
    _dropdown(sheet, lists["FIELD"], f"F{first_empty}:F{last}")
    if first_empty <= last:
        sheet.cell(first_empty, 7, "Empty rows: to correct any other value, give the Collateral ID, the row "
                                   "(and the Neo CIF ID for a BOR row), the field and your value.").font = \
            Font(italic=True, color="FF595959")
    for column, width in zip("ABCDEFGHIJK", (15, 15, 14, 24, 6, 24, 70, 36, 36, 26, 22)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    return len(rows)


def _write_read_me(sheet, batch_name, counts):
    lines = [
        (f"Required details - batch {batch_name}", True),
        ("", False),
        ("The tool could not find every value in CKYC, COLLATERAL and the title reports. Fill this file,", False),
        ("save it (keep its name and place), then run step 2 (2 Apply required details).", False),
        ("", False),
        (f"Borrowers ({counts['borrowers']} rows): one row per borrower that needs a value. A value typed here "
         "is used on every SI of that application.", False),
        (f"Collaterals ({counts['collaterals']} rows): one row per collateral that needs a value.", False),
        (f"Check and correct ({counts['checks']} rows): LMS and title report differences shown side by side, "
         "title reports that may belong to another property, owner names to confirm. Type in 'Your value' "
         "only if the value the tool will use is wrong. Empty rows at the bottom let you correct any other "
         "value.", False),
        ("", False),
        ("Colours", True),
        ("  yellow, empty   needed: type the value", False),
        ("  amber           a default the tool used: leave it to accept it, or type over it", False),
        ("  grey            found by the tool: leave it, or type over it to correct it", False),
        ("  orange          the LMS and the title report disagree (see Check and correct)", False),
        ("  blue            typed for the same collateral or borrower in an earlier batch", False),
        ("  dark grey       does not apply to this borrower or collateral", False),
        ("", False),
        ("Tips", True),
        ("  Dates: YYYY-MM-DD (DD-MM-YYYY is also understood).", False),
        ("  City and district: type the name or the CERSAI code. When the PIN is right, they are worked out "
         "from it, so fill the PIN first.", False),
        ("  Owners: names separated by ;  (e.g. Ramesh Patil; Sunita Patil).", False),
        (f"  To leave an optional value empty in the batch file, type {CLEAR_WORD}.", False),
        ("  Hover over a cell to see why it is needed or where its value came from.", False),
        ("  You can run step 2 as often as you like; each run reads this file again.", False),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        cell = sheet.cell(r, 1, text)
        if bold:
            cell.font = Font(bold=True, size=13 if r == 1 else 11)
    sheet.column_dimensions["A"].width = 140


# --------------------------------------------------------------------------
# Reading
# --------------------------------------------------------------------------
class RequiredDetailsError(Exception):
    """The file cannot be read (missing sheet or renamed header)."""


def _columns(sheet, required) -> dict:
    headers = {as_text(cell.value).upper(): i for i, cell in enumerate(sheet[1]) if as_text(cell.value)}
    missing = [name for name in required if name.upper() not in headers]
    if missing:
        raise RequiredDetailsError(f"sheet '{sheet.title}' is missing the column(s) {', '.join(missing)}; "
                                   "do not rename or delete the header row")
    return {name: headers[name.upper()] for name in required}


def _value(row, index):
    return as_text(row[index]) if index < len(row) else ""


def read_required_details(path, blocks_without_typing, masters):
    """Returns (Overrides, notes). blocks_without_typing: the batch built with
    no typed values, i.e. what the tool finds on its own."""
    book = load_workbook(path, data_only=True)
    for name in ("Borrowers", "Collaterals", "Check and correct"):
        if name not in book.sheetnames:
            raise RequiredDetailsError(f"{path.name} has no sheet '{name}'")
    overrides, notes = Overrides(), []
    borrowers = collect_borrowers(blocks_without_typing)
    by_collateral = {b.collateral_id: b for b in blocks_without_typing}
    match_groups = owner_match_groups(blocks_without_typing)

    # Borrowers
    sheet = book["Borrowers"]
    cols = _columns(sheet, ["Application number", "Neo CIF ID"] + BORROWER_KEYS)
    for number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        app, cif = _value(row, cols["Application number"]), _value(row, cols["Neo CIF ID"])
        if not app and not cif:
            continue
        entry = borrowers.get((app, cif))
        if entry is None:
            notes.append(f"Borrowers row {number}: {app} / {cif} is not a borrower of this batch; ignored")
            continue
        found = borrower_displays(masters, entry)
        for key in BORROWER_KEYS:
            typed = _value(row, cols[key])
            if typed and normalise(key, typed) != normalise(key, found[key]):
                overrides.borrower[(app, cif, key)] = CLEAR_WORD if is_clear(typed) else normalise(key, typed)

    # Collaterals
    sheet = book["Collaterals"]
    cols = _columns(sheet, ["Collateral ID"] + COLLATERAL_KEYS)
    for number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        coll = _value(row, cols["Collateral ID"])
        if not coll:
            continue
        block = by_collateral.get(coll)
        if block is None:
            notes.append(f"Collaterals row {number}: {coll} is not in this batch; ignored")
            continue
        found = collateral_displays(masters, block)
        for key in COLLATERAL_KEYS:
            typed = _value(row, cols[key])
            if typed and normalise(key, typed) != normalise(key, found[key]):
                overrides.collateral[(coll, key)] = CLEAR_WORD if is_clear(typed) else normalise(key, typed)

    # Check and correct (applied last)
    sheet = book["Check and correct"]
    cols = _columns(sheet, ["Collateral ID", "Borrower Neo CIF ID", "Row", "Field", "Your value"])
    known_fields = set(si_row.FIELDS) | set(imm_row.FIELDS) | set(bor_row.FIELDS) | SPECIAL_FIELDS \
        | set(BORROWER_KEYS) | set(COLLATERAL_KEYS)
    for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        typed = _value(row, cols["Your value"])
        if not typed:
            continue
        coll = _value(row, cols["Collateral ID"])
        kind = _value(row, cols["Row"]).upper()
        field = _value(row, cols["Field"]).upper().replace(" ", "_")
        cif = _value(row, cols["Borrower Neo CIF ID"])
        block = by_collateral.get(coll)
        where = f"Check and correct row {number}"
        if block is None:
            notes.append(f"{where}: collateral '{coll}' is not in this batch; ignored")
            continue
        if field not in known_fields:
            notes.append(f"{where}: unknown field '{field}'; ignored")
            continue
        if not kind:
            kind = "BOR" if cif else "IMM"
        value = CLEAR_WORD if is_clear(typed) else normalise(field, typed)
        if field == "PDF_CHECK":
            text = typed.upper()
            value = WRONG_PDF if "WRONG" in text else PDF_OK if "OK" in text else ""
            if not value:
                notes.append(f"{where}: type {PDF_OK} or {WRONG_PDF}; '{typed}' ignored")
                continue
            overrides.collateral[(coll, field)] = value
        elif field in ("OWNER_MATCH", "IS_ASSET_OWNER"):
            if not cif:
                notes.append(f"{where}: give the borrower's Neo CIF ID; ignored")
                continue
            answer = typed.strip().upper()[:1]
            if field == "OWNER_MATCH" and answer not in ("Y", "N"):
                notes.append(f"{where}: type Y or N; '{typed}' ignored")
                continue
            if field == "IS_ASSET_OWNER":
                overrides.per_owner[(coll, cif, field)] = typed
                continue
            # the same owner on other collaterals of the application gets the same answer
            targets = [coll]
            for key, colls in match_groups.items():
                if key[0] == block.application_number and key[1] == cif and coll in colls:
                    targets += [c for c in colls if c not in targets]
            for target in targets:
                overrides.per_owner[(target, cif, field)] = answer
        elif kind == "BOR":
            if not cif or cif not in {b["neo_cif"] for b in block.context.borrowers}:
                notes.append(f"{where}: give the Neo CIF ID of a borrower of {coll}; ignored")
                continue
            short = BORROWER_KEY_OF_FIELD.get(field, field)
            overrides.borrower[(block.application_number, cif, short)] = value
        else:
            overrides.collateral[(coll, field)] = value
    book.close()
    return overrides, notes
